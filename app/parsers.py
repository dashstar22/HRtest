"""简历文件解析：PDF / DOCX / TXT(含 MD) -> 纯文本。"""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_CHARS = 12_000

_NAME_LINE_RE = re.compile(r"(?:姓\s*名|名字)\s*[:：][ \t]*([\u4e00-\u9fa5A-Za-z·]{2,10})")
_FILE_NOISE_RE = re.compile(r"(个人)?简历|resume|cv|副本|final|v\d+", re.IGNORECASE)
_FILE_SPLIT_RE = re.compile(r"[_\-—\s()（）\[\]【】\d.]+")


class ParseError(Exception):
    """解析失败时抛出，带用户可读的中文提示。"""


def _read_pdf(data: bytes) -> str:
    import pdfplumber

    chunks: list[str] = []
    page_count = 0
    image_only_pages = 0

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text = (page.extract_text() or "").strip()
            if text:
                chunks.append(text)
            elif page.images:
                # 这一页只有图片、没有文字层，典型的扫描件特征
                image_only_pages += 1

    text = "\n".join(chunks)
    if len(text) >= 30:
        return text

    if image_only_pages:
        raise ParseError(
            f"这份 PDF 是扫描件：{image_only_pages}/{page_count} 页只有图片、没有文字层，"
            "系统无法直接读取内容。请让候选人提供原始 Word 或文本版 PDF；"
            "或先用 OCR 工具转换后再上传。"
        )
    raise ParseError("PDF 中没有可提取的文字，文件可能是空白或已损坏。")


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    chunks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _read_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="ignore")


def _normalize(text: str) -> str:
    # 统一换行、压缩多余空行与行内空白，保留段落结构供模型定位证据。
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_CHARS]


def guess_name(file_name: str, text: str) -> str:
    match = _NAME_LINE_RE.search(text[:600])
    if match:
        name = re.sub(r"\s+", "", match.group(1))
        if 2 <= len(name) <= 12:
            return name
    stem = Path(file_name).stem
    stem = _FILE_NOISE_RE.sub("", stem)
    parts = [p for p in _FILE_SPLIT_RE.split(stem) if p]
    if parts:
        candidate = max(parts, key=len)
        if 2 <= len(candidate) <= 12:
            return candidate
    return "未命名候选人"


def parse_resume(file_name: str, data: bytes) -> tuple[str, str]:
    """返回 (候选人姓名, 简历正文)。"""
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        if suffix == ".doc":
            raise ParseError("暂不支持旧版 .doc 格式，请另存为 .docx 或 PDF 后重试。")
        raise ParseError(f"不支持的文件类型：{suffix or '未知'}，请上传 PDF / DOCX / TXT。")

    try:
        if suffix == ".pdf":
            raw = _read_pdf(data)
        elif suffix == ".docx":
            raw = _read_docx(data)
        else:
            raw = _read_text(data)
    except ParseError:
        raise
    except Exception as exc:  # 库内部异常统一转成可展示的提示
        raise ParseError(f"文件解析失败：{exc}") from exc

    text = _normalize(raw)
    if len(text) < 30:
        raise ParseError("未能从文件中提取到有效文字（可能是扫描件或图片型 PDF）。")
    return guess_name(file_name, text), text
